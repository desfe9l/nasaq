"""Strict-reader check for files produced by test-render-browser.mjs.
Requires python-docx and python-pptx. No rendered-fidelity claim is made for a
particular Microsoft Office release: this verifies package structure/geometry.
"""
from pathlib import Path
from io import BytesIO
from docx import Document
from pptx import Presentation
from PIL import Image

root = Path('.cache/render-test')
doc = Document(root / 'export.docx')
ppt = Presentation(root / 'export.pptx')
assert len(doc.sections) == 2
anchors = doc.element.xpath('//wp:anchor')
assert len(anchors) == 14
ids = doc.element.xpath('//wp:docPr/@id')
assert len(set(ids)) == len(ids)
assert len(ppt.slides) == 1
assert len(ppt.slides[0].shapes) == 7
assert ppt.slide_width == 3657600  # 101.6mm -> EMU
assert ppt.slide_height == 2743200  # 76.2mm -> EMU
for shape in ppt.slides[0].shapes:
    assert shape.width > 0 and shape.height > 0
    image = Image.open(BytesIO(shape.image.blob))
    assert image.format == 'PNG'
    assert image.mode == 'RGBA'
print('PASS: strict DOCX/PPTX parsers; 14 Word anchors, 7 PowerPoint layers; RGBA PNG media')
